from fastapi import FastAPI, UploadFile, File, Form
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
from dotenv import load_dotenv
import os
import re
import logging
import time
from backend.models import ReviewRequest, SummarizeRequest, SummarizeTextRequest, DialogueRequest, ReadersReviewRequest, WorldBuildingFeedbackRequest, WorldBuildingAutoFillRequest, WBReadersReviewRequest, PanelReviewRequest, WBPanelReviewRequest, PanelAnalyzeRequest, WBPanelAnalyzeRequest, ChatRequest, ChatSummarizeRequest, RewriteRequest, ClassroomExerciseRequest, ClassroomLessonRequest, ClassroomPeerReviewRequest, ClassroomTeacherRequest, ClassroomSynthesisRequest, ClassroomAuthorJudgmentRequest
from backend.mistral_utils import stream_critique_from_mistral, summarize_critiques, summarize_single_critique, summarize_text, stream_from_mistral_small, stream_chat_from_mistral, summarize_chat_messages, get_structured_comments, generate_text_from_mistral
from backend.prompts import REVIEWER_PROMPTS, REVIEWER_NAMES, GENERAL_PROMPT, DIALOGUE_PROMPT, REVIEWER_FIRST_NAMES, READERS_PANEL_PROMPT, WB_READERS_PANEL_PROMPT, WORLD_BUILDING_FEEDBACK_PROMPT, WB_AUTOFILL_PROMPTS, CUSTOM_PANEL_PROMPT, WB_CUSTOM_PANEL_PROMPT, PANEL_ANALYZE_READERS_PROMPT, WB_PANEL_ANALYZE_READERS_PROMPT, CHAT_SINGLE_REVIEWER_PROMPT, CHAT_PANEL_PROMPT, CHAT_CONTEXT_CHAPTER, CHAT_CONTEXT_WB, REWRITE_SINGLE_PROMPT, REWRITE_PANEL_PROMPT, REWRITE_GENERIC_PROMPT, REWRITE_CONTEXT_CHAPTER, REWRITE_CONTEXT_WB, REVIEW_DOCUMENT_PROMPT, REVIEW_DOCUMENT_JSON_PROMPT
from backend.students import STUDENTS, STUDENT_IDS, TEACHER_NAME, TEACHER_PROMPT, get_students_list
from backend.auteurs import AUTHORS, get_author_for_genre
from backend.ai_models import MISTRAL_BEST_MODEL, MISTRAL_MEDIUM_MODEL, ANTHROPIC_BEST_MODEL, ANTHROPIC_MEDIUM_MODEL, OPENAI_BEST_MODEL, OPENAI_MEDIUM_MODEL
from backend.anthropic_utils import stream_critique_from_anthropic
from backend.openai_utils import stream_critique_from_openai
from backend.classroom_prompts import EXERCISE_GENERATION_PROMPT, LESSON_PROMPT, AUTHOR_JUDGMENT_PROMPT, STUDENTS_PANEL_PROMPT, TEACHER_CRITIQUE_PROMPT, TEACHER_MODE_FINAL, TEACHER_MODE_ON_DEMAND, TEACHER_REVISION_NOTE, SYNTHESIS_PROMPT

load_dotenv()

logger = logging.getLogger(__name__)
backend_logger = logging.getLogger("backend")
backend_logger.setLevel(logging.INFO)
if not backend_logger.handlers:
    trace_handler = logging.StreamHandler()
    trace_handler.setFormatter(logging.Formatter('%(asctime)s %(levelname)s %(message)s'))
    backend_logger.addHandler(trace_handler)
backend_logger.propagate = False

app = FastAPI()


@app.middleware("http")
async def log_api_requests(request, call_next):
    started_at = time.perf_counter()
    method = request.method
    path = request.url.path
    logger.info("[API] request_started method=%s path=%s", method, path)
    try:
        response = await call_next(request)
    except Exception:
        duration_ms = (time.perf_counter() - started_at) * 1000
        logger.exception("[API] request_failed method=%s path=%s duration_ms=%.1f", method, path, duration_ms)
        raise
    duration_ms = (time.perf_counter() - started_at) * 1000
    logger.info(
        "[API] response_sent method=%s path=%s status=%d duration_ms=%.1f",
        method,
        path,
        response.status_code,
        duration_ms,
    )
    return response

# CORS middleware pour permettre les requêtes depuis le frontend
app.add_middleware(
    CORSMiddleware,
    allow_origins=["http://localhost:3000", "http://127.0.0.1:3000"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

@app.get("/")
async def root():
    return {"message": "FastAPI backend is running!"}

@app.get("/ping")
async def ping():
    return {"status": "ok"}

@app.get("/config")
async def get_config():
    return {
        "mistral_api_key_set": bool(os.getenv("MISTRAL_API_KEY")),
        "anthropic_api_key_set": bool(os.getenv("ANTHROPIC_API_KEY")),
        "openai_api_key_set": bool(os.getenv("OPENAI_API_KEY")),
    }


# Choisit le meilleur fournisseur dispo (clé API présente), avec repli sur Mistral en dernier recours.
def _pick_provider(preferred: list[tuple[str, callable, str]]):
    for env_var, stream_fn, model in preferred:
        if os.getenv(env_var):
            provider = env_var.removesuffix("_API_KEY").lower()
            logger.info("[AI] provider_selected provider=%s model=%s", provider, model)
            return stream_fn, model, provider
    fallback = preferred[-1]
    provider = fallback[0].removesuffix("_API_KEY").lower()
    logger.info("[AI] provider_selected provider=%s model=%s fallback=true", provider, fallback[2])
    return fallback[1], fallback[2], provider


async def _trace_ai_response(stream, provider: str, model: str):
    response_parts: list[str] = []
    logger.info("[AI] request_started provider=%s model=%s", provider, model)
    try:
        async for chunk in stream:
            response_parts.append(chunk)
            yield chunk
    finally:
        response = ''.join(response_parts)
        logger.info(
            "[AI] response_received provider=%s model=%s chars=%d preview=%r",
            provider,
            model,
            len(response),
            response[:500],
        )

@app.get("/reviewers")
async def get_reviewers():
    return [{"id": k, "name": REVIEWER_NAMES.get(k, k)} for k in REVIEWER_PROMPTS.keys()]

@app.post("/review")
async def review_text(request: ReviewRequest):
    if request.reviewer not in REVIEWER_PROMPTS:
        async def error_stream():
            yield "Critique inconnu. Veuillez choisir parmi les critiques disponibles."
        return StreamingResponse(error_stream(), media_type="text/plain")

    reviewer_prompt = REVIEWER_PROMPTS[request.reviewer]
    
    # Build full prompt: general + context + writer profile + previous critiques + reviewer personality
    parts = [GENERAL_PROMPT]
    
    if request.book_summary:
        parts.append(
            f"\n\nRÉSUMÉ DE L'OUVRAGE (pour que tu comprennes le contexte global) : {request.book_summary}"
        )
    
    if request.chapter_summary:
        parts.append(
            f"\n\nRÉSUMÉ DU CHAPITRE (ce que l'auteur veut écrire dans ce chapitre) : {request.chapter_summary}"
        )
    
    if request.writer_profile:
        parts.append(
            f"\n\nProfil de l'auteur que tu critiques (adapte tes conseils en conséquence) : {request.writer_profile}"
        )
    
    if request.previous_critiques:
        # Summarize past critiques via mistral-small-latest to save context
        try:
            summary = await summarize_critiques(request.previous_critiques)
        except Exception:
            summary = "\n".join(f"- {pc.get('reviewer', '?')}: {pc.get('critique', '')[:200]}" for pc in request.previous_critiques)
        
        if request.text_changed:
            parts.append(
                f"\n\nSYNTHÈSE DES CRITIQUES PRÉCÉDENTES (l'auteur a modifié son texte depuis, évalue sa progression) :\n{summary}\n"
                "\nCompare le texte actuel avec ces retours précédents. Note les améliorations et ce qui reste à travailler.\n"
            )
        else:
            parts.append(
                f"\n\nSYNTHÈSE DES CRITIQUES PRÉCÉDENTES (le texte n'a pas changé, l'auteur demande un regard différent) :\n{summary}\n"
                "\nNe répète pas les mêmes remarques. Apporte un regard neuf, complémentaire, en évitant de redire ce qui a déjà été dit.\n"
            )
    
    parts.append(f"\n\n{reviewer_prompt}")
    full_prompt = "".join(parts)
    
    return StreamingResponse(
        stream_critique_from_mistral(request.text, full_prompt),
        media_type="text/plain"
    )


@app.post("/summarize-critique")
async def summarize_critique_endpoint(request: SummarizeRequest):
    try:
        summary = await summarize_single_critique(request.critique)
        return {"summary": summary}
    except Exception as e:
        return {"summary": "", "error": str(e)}


@app.post("/summarize-text")
async def summarize_text_endpoint(request: SummarizeTextRequest):
    try:
        summary = await summarize_text(request.text)
        return {"summary": summary}
    except Exception as e:
        return {"summary": "", "error": str(e)}


@app.post("/review-dialogue")
async def review_dialogue(request: DialogueRequest):
    import random
    
    # Pick N random reviewers
    all_ids = list(REVIEWER_PROMPTS.keys())
    n = min(request.num_authors, len(all_ids))
    chosen = random.sample(all_ids, n)
    
    # Build authors list and personalities
    authors_list = "\n".join(
        f"- {REVIEWER_FIRST_NAMES.get(rid, rid)} ({REVIEWER_NAMES.get(rid, rid)})"
        for rid in chosen
    )
    authors_personalities = "\n\n".join(
        f"{REVIEWER_FIRST_NAMES.get(rid, rid)} :\n{REVIEWER_PROMPTS[rid]}"
        for rid in chosen
    )
    
    # Build the dialogue prompt
    dialogue_system = DIALOGUE_PROMPT.format(
        nb_authors=n,
        authors_list=authors_list,
        authors_personalities=authors_personalities
    )
    
    # Add context
    parts = [dialogue_system]
    
    if request.book_summary:
        parts.append(f"\n\nRÉSUMÉ DE L'OUVRAGE : {request.book_summary}")
    if request.chapter_summary:
        parts.append(f"\n\nRÉSUMÉ DU CHAPITRE : {request.chapter_summary}")
    if request.writer_profile:
        parts.append(f"\n\nProfil de l'auteur : {request.writer_profile}")
    
    full_prompt = "".join(parts)
    chosen_names = [REVIEWER_FIRST_NAMES.get(rid, rid) for rid in chosen]
    
    async def stream_with_header():
        yield f"*{', '.join(chosen_names)} s'installent autour de votre texte…*\n\n---\n\n"
        async for chunk in stream_critique_from_mistral(request.text, full_prompt):
            yield chunk
    
    return StreamingResponse(
        stream_with_header(),
        media_type="text/plain",
        headers={"X-Dialogue-Authors": ",".join(chosen)}
    )


@app.post("/review-readers")
async def review_readers(request: ReadersReviewRequest):
    """Panel de lecteurs aléatoires qui donnent leur avis."""
    readers_system = READERS_PANEL_PROMPT.format(nb_readers=request.num_readers)
    
    parts = [readers_system]
    
    if request.book_summary:
        parts.append(f"\n\nRÉSUMÉ DE L'OUVRAGE : {request.book_summary}")
    if request.chapter_summary:
        parts.append(f"\n\nRÉSUMÉ DU CHAPITRE : {request.chapter_summary}")
    if request.writer_profile:
        parts.append(f"\n\nProfil de l'auteur : {request.writer_profile}")
    
    full_prompt = "".join(parts)
    
    return StreamingResponse(
        stream_critique_from_mistral(request.text, full_prompt),
        media_type="text/plain"
    )


@app.post("/world-building/feedback")
async def world_building_feedback(request: WorldBuildingFeedbackRequest):
    prompt_text = WORLD_BUILDING_FEEDBACK_PROMPT.format(
        category=request.category,
        title=request.entry_title,
    )
    
    parts = [prompt_text]
    
    # Add reviewer personality if specified
    if request.reviewer and request.reviewer in REVIEWER_PROMPTS:
        reviewer_name = REVIEWER_NAMES.get(request.reviewer, request.reviewer)
        parts.append(
            f"\n\nATTENTION — Tu ne parles pas en tant que consultant générique. "
            f"Tu es {reviewer_name}. Adopte complètement sa personnalité, son style, sa voix, ses tics de langage. "
            f"Voici ta personnalité :\n{REVIEWER_PROMPTS[request.reviewer]}\n\n"
            f"Critique cet élément de world building COMME {reviewer_name} le ferait. "
            f"Garde la structure demandée (Ce qui fonctionne bien / Points à améliorer / Suggestions / Cohérence) "
            f"mais écris avec ta voix de {reviewer_name}."
        )
    
    if request.book_summary:
        parts.append(
            f"\n\nRÉSUMÉ DE L'OUVRAGE (contexte global) : {request.book_summary}"
        )
    
    if request.all_entries_context:
        parts.append(
            f"\n\nTOUT LE WORLD BUILDING DE L'UNIVERS (utilise TOUT pour ton analyse de cohérence) :\n{request.all_entries_context}"
        )
    
    full_prompt = "".join(parts)
    
    return StreamingResponse(
        stream_critique_from_mistral(request.entry_content, full_prompt),
        media_type="text/plain"
    )


@app.post("/world-building/readers")
async def world_building_readers(request: WBReadersReviewRequest):
    """Panel de lecteurs aléatoires pour le world building."""
    readers_system = WB_READERS_PANEL_PROMPT.format(
        nb_readers=request.num_readers,
        category=request.category,
        title=request.entry_title,
    )
    
    parts = [readers_system]
    
    if request.book_summary:
        parts.append(f"\n\nRÉSUMÉ DE L'OUVRAGE : {request.book_summary}")
    
    if request.all_entries_context:
        parts.append(f"\n\nAUTRES ÉLÉMENTS DE L'UNIVERS :\n{request.all_entries_context}")
    
    full_prompt = "".join(parts)
    
    return StreamingResponse(
        stream_critique_from_mistral(request.entry_content, full_prompt),
        media_type="text/plain"
    )


@app.post("/review-panel")
async def review_custom_panel(request: PanelReviewRequest):
    """Panel personnalisé d'auteurs choisis par l'utilisateur."""
    # Filter valid reviewer IDs
    chosen = [rid for rid in request.reviewer_ids if rid in REVIEWER_PROMPTS]
    if not chosen:
        async def error_stream():
            yield "Aucun auteur valide dans votre panel. Configurez votre panel dans les param\u00e8tres."
        return StreamingResponse(error_stream(), media_type="text/plain")
    
    n = len(chosen)
    authors_list = "\n".join(
        f"- {REVIEWER_FIRST_NAMES.get(rid, rid)} ({REVIEWER_NAMES.get(rid, rid)})"
        for rid in chosen
    )
    authors_personalities = "\n\n".join(
        f"{REVIEWER_FIRST_NAMES.get(rid, rid)} :\n{REVIEWER_PROMPTS[rid]}"
        for rid in chosen
    )
    
    panel_system = CUSTOM_PANEL_PROMPT.format(
        nb_authors=n,
        authors_list=authors_list,
        authors_personalities=authors_personalities
    )
    
    parts = [panel_system]
    if request.book_summary:
        parts.append(f"\n\nR\u00c9SUM\u00c9 DE L'OUVRAGE : {request.book_summary}")
    if request.chapter_summary:
        parts.append(f"\n\nR\u00c9SUM\u00c9 DU CHAPITRE : {request.chapter_summary}")
    if request.writer_profile:
        parts.append(f"\n\nProfil de l'auteur : {request.writer_profile}")
    
    full_prompt = "".join(parts)
    chosen_names = [REVIEWER_FIRST_NAMES.get(rid, rid) for rid in chosen]
    
    async def stream_with_header():
        yield f"*\u2b50 Votre panel : {', '.join(chosen_names)} prennent place autour de votre texte\u2026*\n\n---\n\n"
        async for chunk in stream_critique_from_mistral(request.text, full_prompt):
            yield chunk
    
    return StreamingResponse(
        stream_with_header(),
        media_type="text/plain",
        headers={"X-Panel-Authors": ",".join(chosen)}
    )


@app.post("/world-building/panel")
async def world_building_custom_panel(request: WBPanelReviewRequest):
    """Panel personnalisé d'auteurs pour le world building."""
    chosen = [rid for rid in request.reviewer_ids if rid in REVIEWER_PROMPTS]
    if not chosen:
        async def error_stream():
            yield "Aucun auteur valide dans votre panel. Configurez votre panel dans les param\u00e8tres."
        return StreamingResponse(error_stream(), media_type="text/plain")
    
    n = len(chosen)
    authors_list = "\n".join(
        f"- {REVIEWER_FIRST_NAMES.get(rid, rid)} ({REVIEWER_NAMES.get(rid, rid)})"
        for rid in chosen
    )
    authors_personalities = "\n\n".join(
        f"{REVIEWER_FIRST_NAMES.get(rid, rid)} :\n{REVIEWER_PROMPTS[rid]}"
        for rid in chosen
    )
    
    panel_system = WB_CUSTOM_PANEL_PROMPT.format(
        nb_authors=n,
        authors_list=authors_list,
        authors_personalities=authors_personalities,
        category=request.category,
        title=request.entry_title
    )
    
    parts = [panel_system]
    if request.book_summary:
        parts.append(f"\n\nR\u00c9SUM\u00c9 DE L'OUVRAGE : {request.book_summary}")
    if request.all_entries_context:
        parts.append(f"\n\nAUTRES \u00c9L\u00c9MENTS DE L'UNIVERS :\n{request.all_entries_context}")
    
    full_prompt = "".join(parts)
    chosen_names = [REVIEWER_FIRST_NAMES.get(rid, rid) for rid in chosen]
    
    async def stream_with_header():
        yield f"*\u2b50 Votre panel : {', '.join(chosen_names)} examinent votre univers\u2026*\n\n---\n\n"
        async for chunk in stream_critique_from_mistral(request.entry_content, full_prompt):
            yield chunk
    
    return StreamingResponse(
        stream_with_header(),
        media_type="text/plain",
        headers={"X-Panel-Authors": ",".join(chosen)}
    )


@app.post("/review-panel/analyze")
async def review_panel_analyze(request: PanelAnalyzeRequest):
    """Le panel analyse les retours des lecteurs et propose des modifications."""
    chosen = [rid for rid in request.reviewer_ids if rid in REVIEWER_PROMPTS]
    if not chosen:
        async def error_stream():
            yield "Aucun auteur valide dans votre panel. Configurez votre panel dans les paramètres."
        return StreamingResponse(error_stream(), media_type="text/plain")

    n = len(chosen)
    authors_list = "\n".join(
        f"- {REVIEWER_FIRST_NAMES.get(rid, rid)} ({REVIEWER_NAMES.get(rid, rid)})"
        for rid in chosen
    )
    authors_personalities = "\n\n".join(
        f"{REVIEWER_FIRST_NAMES.get(rid, rid)} :\n{REVIEWER_PROMPTS[rid]}"
        for rid in chosen
    )

    panel_system = PANEL_ANALYZE_READERS_PROMPT.format(
        nb_authors=n,
        authors_list=authors_list,
        authors_personalities=authors_personalities,
        reader_feedback=request.reader_feedback
    )

    parts = [panel_system]
    if request.book_summary:
        parts.append(f"\n\nRÉSUMÉ DE L'OUVRAGE : {request.book_summary}")
    if request.chapter_summary:
        parts.append(f"\n\nRÉSUMÉ DU CHAPITRE : {request.chapter_summary}")
    if request.writer_profile:
        parts.append(f"\n\nProfil de l'auteur : {request.writer_profile}")

    full_prompt = "".join(parts)
    chosen_names = [REVIEWER_FIRST_NAMES.get(rid, rid) for rid in chosen]

    async def stream_with_header():
        yield f"*🔍 Votre panel : {', '.join(chosen_names)} analysent les retours des lecteurs…*\n\n---\n\n"
        async for chunk in stream_critique_from_mistral(request.text, full_prompt):
            yield chunk

    return StreamingResponse(
        stream_with_header(),
        media_type="text/plain",
        headers={"X-Panel-Authors": ",".join(chosen)}
    )


@app.post("/world-building/panel/analyze")
async def wb_panel_analyze(request: WBPanelAnalyzeRequest):
    """Le panel analyse les retours des lecteurs sur le world building."""
    chosen = [rid for rid in request.reviewer_ids if rid in REVIEWER_PROMPTS]
    if not chosen:
        async def error_stream():
            yield "Aucun auteur valide dans votre panel. Configurez votre panel dans les paramètres."
        return StreamingResponse(error_stream(), media_type="text/plain")

    n = len(chosen)
    authors_list = "\n".join(
        f"- {REVIEWER_FIRST_NAMES.get(rid, rid)} ({REVIEWER_NAMES.get(rid, rid)})"
        for rid in chosen
    )
    authors_personalities = "\n\n".join(
        f"{REVIEWER_FIRST_NAMES.get(rid, rid)} :\n{REVIEWER_PROMPTS[rid]}"
        for rid in chosen
    )

    panel_system = WB_PANEL_ANALYZE_READERS_PROMPT.format(
        nb_authors=n,
        authors_list=authors_list,
        authors_personalities=authors_personalities,
        category=request.category,
        title=request.entry_title,
        reader_feedback=request.reader_feedback
    )

    parts = [panel_system]
    if request.book_summary:
        parts.append(f"\n\nRÉSUMÉ DE L'OUVRAGE : {request.book_summary}")
    if request.all_entries_context:
        parts.append(f"\n\nAUTRES ÉLÉMENTS DE L'UNIVERS :\n{request.all_entries_context}")

    full_prompt = "".join(parts)
    chosen_names = [REVIEWER_FIRST_NAMES.get(rid, rid) for rid in chosen]

    async def stream_with_header():
        yield f"*🔍 Votre panel : {', '.join(chosen_names)} analysent les retours sur « {request.entry_title} »…*\n\n---\n\n"
        async for chunk in stream_critique_from_mistral(request.entry_content, full_prompt):
            yield chunk

    return StreamingResponse(
        stream_with_header(),
        media_type="text/plain",
        headers={"X-Panel-Authors": ",".join(chosen)}
    )


@app.post("/world-building/autofill")
async def world_building_autofill(request: WorldBuildingAutoFillRequest):
    # Find the right prompt for this category
    category_key = request.category.lower()
    # Try to match category key from the prompts dict
    prompt = WB_AUTOFILL_PROMPTS.get(category_key, WB_AUTOFILL_PROMPTS["autre"])
    
    # Build user content
    user_parts = [f"Titre de l'élément : « {request.entry_title} »\n"]
    user_parts.append(f"Catégorie : {request.category}\n")
    
    if request.entry_content and request.entry_content.strip():
        user_parts.append(
            f"\nCONTENU EXISTANT (à développer et enrichir, ne pas remplacer) :\n"
            f"{request.entry_content}\n\n"
            "Développe, enrichis et complète ce qui est écrit ci-dessus. "
            "Garde le ton et le style de l'auteur. Ajoute des détails, de la profondeur, des nuances."
        )
    else:
        user_parts.append(
            "\nAucun contenu existant. Invente tout à partir du titre et du contexte de l'univers."
        )
    
    if request.book_summary:
        user_parts.append(f"\n\nRÉSUMÉ DE L'OUVRAGE : {request.book_summary}")
    
    if request.all_entries_context:
        user_parts.append(
            f"\n\nAUTRES ÉLÉMENTS DE L'UNIVERS (reste cohérent avec eux) :\n{request.all_entries_context}"
        )
    
    user_content = "".join(user_parts)
    
    return StreamingResponse(
        stream_from_mistral_small(prompt, user_content),
        media_type="text/plain"
    )


@app.post("/chat")
async def chat_with_reviewer(request: ChatRequest):
    """Chat multi-tour avec un auteur ou un panel, avec prompt caching Mistral."""
    
    # Build system prompt based on who we're chatting with
    if request.reviewer_ids and len(request.reviewer_ids) > 0:
        # Panel mode
        chosen = [rid for rid in request.reviewer_ids if rid in REVIEWER_PROMPTS]
        if not chosen:
            async def error_stream():
                yield "Aucun auteur valide dans votre panel."
            return StreamingResponse(error_stream(), media_type="text/plain")
        
        n = len(chosen)
        authors_list = "\n".join(
            f"- {REVIEWER_FIRST_NAMES.get(rid, rid)} ({REVIEWER_NAMES.get(rid, rid)})"
            for rid in chosen
        )
        authors_personalities = "\n\n".join(
            f"{REVIEWER_FIRST_NAMES.get(rid, rid)} :\n{REVIEWER_PROMPTS[rid]}"
            for rid in chosen
        )
        system_prompt = CHAT_PANEL_PROMPT.format(
            nb_authors=n,
            authors_list=authors_list,
            authors_personalities=authors_personalities
        )
    elif request.reviewer_id and request.reviewer_id in REVIEWER_PROMPTS:
        # Single reviewer mode
        rid = request.reviewer_id
        system_prompt = CHAT_SINGLE_REVIEWER_PROMPT.format(
            reviewer_name=REVIEWER_NAMES.get(rid, rid),
            reviewer_first_name=REVIEWER_FIRST_NAMES.get(rid, rid),
            reviewer_personality=REVIEWER_PROMPTS[rid]
        )
    else:
        # Generic consultant
        system_prompt = (
            "Tu es un consultant littéraire expert, bienveillant et constructif. "
            "Tu discutes avec l'auteur de son texte. "
            "Quand il te demande des propositions de réécriture, tu proposes des passages concrets "
            "délimités par des balises ```suggestion\n...\n```.\n"
            "Tu cites le passage original en italique avant ta proposition.\n"
            "Réponds toujours en français, sois concis mais substantiel (200-400 mots)."
        )
    
    # Add context type
    if request.context_type == "world_building":
        system_prompt += CHAT_CONTEXT_WB.format(
            category=request.category or "Général",
            title=request.entry_title or "Sans titre"
        )
    else:
        system_prompt += CHAT_CONTEXT_CHAPTER
    
    # Add meta-context
    if request.book_summary:
        system_prompt += f"\nRESUME DE L'OUVRAGE : {request.book_summary}"
    if request.chapter_summary:
        system_prompt += f"\nRESUME DU CHAPITRE : {request.chapter_summary}"
    if request.writer_profile:
        system_prompt += f"\nPROFIL DE L'AUTEUR : {request.writer_profile}"
    if request.all_entries_context:
        system_prompt += f"\nAUTRES ÉLÉMENTS DE L'UNIVERS :\n{request.all_entries_context}"
    
    # Build messages array for Mistral (prefix stays stable for caching)
    messages = [{"role": "system", "content": system_prompt}]
    
    # First user message: the original text
    if request.context_type == "world_building":
        messages.append({"role": "user", "content": (
            f"Voici l'élément de world building que tu viens de critiquer :\n\n"
            f"[DÉBUT DU CONTENU]\n{request.text}\n[FIN DU CONTENU]"
        )})
    else:
        messages.append({"role": "user", "content": (
            f"Voici le texte que tu viens de critiquer :\n\n"
            f"[DÉBUT DU TEXTE À CRITIQUER]\n{request.text}\n[FIN DU TEXTE À CRITIQUER]"
        )})
    
    # First assistant message: the initial feedback/critique
    messages.append({"role": "assistant", "content": request.initial_feedback})
    
    # Inject accumulated summaries of past exchanges (if any)
    # The prefix above (system + text + critique) stays identical → Mistral caches it
    if request.chat_summaries and len(request.chat_summaries) > 0:
        summaries_text = "\n\n---\n\n".join(
            f"**Résumé échange {i+1} :** {s}"
            for i, s in enumerate(request.chat_summaries)
        )
        messages.append({"role": "user", "content": (
            f"[Voici un résumé de nos échanges précédents pour contexte :]\n\n{summaries_text}"
        )})
        messages.append({"role": "assistant", "content": (
            "Compris, j'ai bien le contexte de nos échanges précédents. Continuons."
        )})
    
    # Recent conversation history (only unsummarized messages sent by frontend)
    for msg in request.messages:
        messages.append({"role": msg.role, "content": msg.content})
    
    return StreamingResponse(
        stream_chat_from_mistral(messages),
        media_type="text/plain"
    )


@app.post("/rewrite")
async def rewrite_text(request: RewriteRequest):
    """Stream a rewritten version of the text from a reviewer."""
    import random

    # Build the system prompt based on reviewer type
    if request.reviewer_ids and len(request.reviewer_ids) > 0:
        chosen = [rid for rid in request.reviewer_ids if rid in REVIEWER_PROMPTS]
        if not chosen:
            chosen = random.sample(list(REVIEWER_PROMPTS.keys()), min(3, len(REVIEWER_PROMPTS)))
        n = len(chosen)
        authors_list = ", ".join(REVIEWER_NAMES.get(rid, rid) for rid in chosen)
        authors_personalities = "\n\n".join(
            f"{REVIEWER_FIRST_NAMES.get(rid, rid)} :\n{REVIEWER_PROMPTS[rid]}"
            for rid in chosen
        )
        system_prompt = REWRITE_PANEL_PROMPT.format(
            nb_authors=n,
            authors_list=authors_list,
            authors_personalities=authors_personalities
        )
    elif request.reviewer_id and request.reviewer_id in REVIEWER_PROMPTS:
        rid = request.reviewer_id
        system_prompt = REWRITE_SINGLE_PROMPT.format(
            reviewer_name=REVIEWER_NAMES.get(rid, rid),
            reviewer_personality=REVIEWER_PROMPTS[rid]
        )
    else:
        system_prompt = REWRITE_GENERIC_PROMPT

    # Add context type
    if request.context_type == "world_building":
        system_prompt += REWRITE_CONTEXT_WB.format(
            category=request.category or "Général",
            title=request.entry_title or "Sans titre"
        )
    else:
        system_prompt += REWRITE_CONTEXT_CHAPTER

    # Add meta-context
    if request.book_summary:
        system_prompt += f"\nRÉSUMÉ DE L'OUVRAGE : {request.book_summary}"
    if request.chapter_summary:
        system_prompt += f"\nRÉSUMÉ DU CHAPITRE : {request.chapter_summary}"
    if request.writer_profile:
        system_prompt += f"\nPROFIL DE L'AUTEUR : {request.writer_profile}"
    if request.all_entries_context:
        system_prompt += f"\nAUTRES ÉLÉMENTS DE L'UNIVERS :\n{request.all_entries_context}"

    # Build user message
    user_content = f"[DÉBUT DU TEXTE À RÉÉCRIRE]\n{request.text}\n[FIN DU TEXTE À RÉÉCRIRE]"
    if request.instructions:
        user_content += f"\n\nINSTRUCTIONS SPÉCIFIQUES DE L'AUTEUR :\n{request.instructions}"

    messages = [
        {"role": "system", "content": system_prompt},
        {"role": "user", "content": user_content}
    ]

    return StreamingResponse(
        stream_chat_from_mistral(messages),
        media_type="text/plain"
    )


@app.post("/chat/summarize")
async def summarize_chat(request: ChatSummarizeRequest):
    """Résume un batch de messages de chat via mistral-small."""
    msgs = [{"role": m.role, "content": m.content} for m in request.messages]
    summary = await summarize_chat_messages(msgs)
    return {"summary": summary}


@app.post("/import-document")
async def import_document(file: UploadFile = File(...)):
    """Importe un .docx ou .odt et retourne {title, chapters}."""
    from backend.document_import import parse_docx, parse_odt

    filename = file.filename or "document"
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""

    if ext not in ("docx", "doc", "odt"):
        return {"error": "Format non supporté. Utilisez .docx ou .odt"}

    file_bytes = await file.read()

    try:
        if ext in ("docx", "doc"):
            result = parse_docx(file_bytes)
        else:
            result = parse_odt(file_bytes)
    except Exception as e:
        return {"error": f"Erreur de lecture du document : {e}"}

    if not result.get("title"):
        result["title"] = filename.rsplit(".", 1)[0]

    return result


@app.post("/review-document")
async def review_document(file: UploadFile = File(...), reviewer: str = Form("prof_ecriture"), context: str = Form("")):
    """Importe un document et produit un .docx commenté avec des annotations Word."""
    from backend.document_import import extract_structured_text
    from backend.docx_comments import build_commented_docx
    import json as json_module
    import traceback

    if reviewer not in REVIEWER_PROMPTS:
        return {"error": "Reviewer inconnu"}

    filename = file.filename or "document"
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext not in ("docx", "doc", "odt"):
        return {"error": "Format non supporté. Utilisez .docx ou .odt"}

    file_bytes = await file.read()
    print(f"[review-document] Fichier: {filename}, ext: {ext}, taille: {len(file_bytes)} octets")

    try:
        result = extract_structured_text(file_bytes, ext)
    except Exception as e:
        traceback.print_exc()
        return {"error": f"Erreur de lecture du document : {e}"}

    doc_title = result.get("title") or filename.rsplit(".", 1)[0]
    sections = result.get("sections", [])
    print(f"[review-document] Titre: {doc_title}, sections: {len(sections)}")

    if not sections:
        return {"error": "Aucun contenu trouvé dans le document"}

    # Construire le texte pour l'IA
    doc_text_parts = [f"TITRE DU DOCUMENT : {doc_title}\n"]
    for sec in sections:
        heading = sec["heading"]
        level = sec.get("level", 1)
        hashes = "#" * min(level, 4)
        doc_text_parts.append(f"\n{hashes} {heading}\n")
        for para in sec["paragraphs"]:
            doc_text_parts.append(para)

    doc_text = "\n".join(doc_text_parts)
    print(f"[review-document] Texte pour IA: {len(doc_text)} caractères")

    # Construire le prompt système pour obtenir du JSON
    reviewer_prompt = REVIEWER_PROMPTS[reviewer]
    reviewer_name = REVIEWER_FIRST_NAMES.get(reviewer, "Le Critique")
    system_prompt = (
        REVIEW_DOCUMENT_JSON_PROMPT
        + f"\n\nTA PERSONA :\n{reviewer_prompt}\n"
    )
    if context:
        system_prompt += f"\nCONTEXTE FOURNI PAR L'AUTEUR (ce qu'est ce document, à qui il s'adresse, ce qu'il en attend) :\n{context}\n"

    # Appel non-streaming pour obtenir le JSON des commentaires
    print(f"[review-document] Appel IA en cours...")
    try:
        raw_response = await get_structured_comments(doc_text, system_prompt)
    except Exception as e:
        traceback.print_exc()
        return {"error": f"Erreur lors de l'analyse IA : {e}"}

    print(f"[review-document] Réponse IA reçue: {len(raw_response)} caractères")

    # Parser le JSON
    try:
        # Nettoyer la réponse (enlever markdown fences si présents)
        cleaned = raw_response.strip()
        if cleaned.startswith("```"):
            first_newline = cleaned.find("\n")
            if first_newline != -1:
                cleaned = cleaned[first_newline + 1:]
            else:
                cleaned = cleaned[3:]
        if cleaned.endswith("```"):
            cleaned = cleaned[:-3]
        cleaned = cleaned.strip()

        def fix_json_newlines(s: str) -> str:
            """Parcourt le JSON caractère par caractère et remplace les vrais
            sauts de ligne et caractères de contrôle à l'intérieur des chaînes."""
            result = []
            in_string = False
            escape = False
            for ch in s:
                if escape:
                    result.append(ch)
                    escape = False
                    continue
                if ch == '\\' and in_string:
                    result.append(ch)
                    escape = True
                    continue
                if ch == '"':
                    in_string = not in_string
                    result.append(ch)
                    continue
                if in_string and ord(ch) < 32:
                    # Remplacer tous les caractères de contrôle (newline, tab, etc.) par un espace
                    if ch not in ('\n', '\r', '\t'):
                        continue  # supprimer les autres caractères de contrôle
                    result.append(' ')
                    continue
                result.append(ch)
            return ''.join(result)

        # Tenter de parser tel quel, sinon réparer les newlines
        try:
            comments = json_module.loads(cleaned)
        except json_module.JSONDecodeError:
            fixed = fix_json_newlines(cleaned)
            comments = json_module.loads(fixed)

        if not isinstance(comments, list):
            comments = [comments]
    except json_module.JSONDecodeError as e:
        print(f"[review-document] ERREUR JSON: {e}")
        print(f"[review-document] Réponse brute: {raw_response[:500]}")
        return {"error": f"Erreur de parsing JSON de l'IA : {e}\nRéponse brute : {raw_response[:500]}"}

    print(f"[review-document] {len(comments)} commentaires parsés, construction du .docx...")

    # Construire le .docx commenté
    try:
        docx_bytes = build_commented_docx(
            original_bytes=file_bytes,
            comments=comments,
            author=reviewer_name,
            ext=ext,
        )
    except Exception as e:
        traceback.print_exc()
        return {"error": f"Erreur lors de la construction du document commenté : {e}"}

    print(f"[review-document] .docx généré: {len(docx_bytes)} octets")

    # Retourner le .docx
    output_name = f"commenté-{filename.rsplit('.', 1)[0]}.docx"
    from fastapi.responses import Response
    return Response(
        content=docx_bytes,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={
            "Content-Disposition": f'attachment; filename="{output_name}"',
        },
    )


@app.post("/import-document/debug")
async def import_document_debug(file: UploadFile = File(...)):
    """Debug : montre les styles et niveaux détectés pour chaque paragraphe."""
    from docx import Document
    from docx.text.paragraph import Paragraph
    import io

    filename = file.filename or "document"
    file_bytes = await file.read()
    doc = Document(io.BytesIO(file_bytes))

    paras = []
    for i, para in enumerate(doc.paragraphs):
        text = para.text.strip()
        if not text:
            continue
        style_name = para.style.name if para.style else ""
        style_id = para.style.style_id if para.style else ""
        from backend.document_import import _detect_heading_level_docx
        level = _detect_heading_level_docx(para, style_name)
        paras.append({
            "index": i,
            "style_name": style_name,
            "style_id": style_id,
            "heading_level": level,
            "text_preview": text[:120],
        })

    return {"filename": filename, "total_paragraphs": len(doc.paragraphs), "detected": paras[:100]}


# ==========================================
# CLASSE D'ÉCRITURE CRÉATIVE ("J'apprends")
# ==========================================

@app.get("/classroom/students")
async def classroom_students():
    return {"teacher": {"name": TEACHER_NAME}, "students": get_students_list()}


_CLASSROOM_EXPERIENCE_CONTEXTS = {
    "grand_debutant": (
        "Grand débutant", 1,
        "Privilégie une seule compétence à la fois, des consignes très claires et des contraintes simples. "
        "Explique les termes techniques et valorise avant tout l'élan, la compréhension de la consigne et les progrès visibles."
    ),
    "debutant": (
        "Débutant", 2,
        "Propose des exercices accessibles avec une technique identifiable et peu de contraintes cumulées. "
        "Les retours doivent rester très concrets, pédagogiques et proposer une priorité d'amélioration à la fois."
    ),
    "intermediaire": (
        "Intermédiaire", 3,
        "Propose des exercices qui demandent de maîtriser une technique narrative ou stylistique. "
        "Les retours peuvent employer le vocabulaire d'atelier en l'explicitant et doivent relever les choix de structure, de voix et de rythme."
    ),
    "avance": (
        "Avancé", 4,
        "Propose des exercices ambitieux à contraintes croisées, qui demandent une intention littéraire consciente. "
        "Les retours doivent être précis et exigeants sur la maîtrise du style, de la structure et des effets produits."
    ),
    "ecrivain_publie": (
        "Écrivain publié", 5,
        "Propose des exercices de haut niveau, formellement ambitieux et proches d'un travail de publication. "
        "Les retours doivent être francs, rigoureux et éditoriaux, en questionnant la singularité, la cohérence d'ensemble et l'impact sur le lecteur."
    ),
}


def _build_experience_context(experience_level: str) -> str:
    """Décrit le niveau de l'élève pour adapter chaque interaction de l'atelier."""
    label, difficulty, guidance = _CLASSROOM_EXPERIENCE_CONTEXTS[experience_level]
    return (
        f"NIVEAU D'EXPÉRIENCE DE L'ÉLÈVE : {label} (difficulté cible {difficulty}/5).\n"
        f"ADAPTATION PÉDAGOGIQUE OBLIGATOIRE : {guidance}"
    )


def _build_student_record(past_syntheses: list[dict] | None) -> str:
    """Formate le carnet de suivi de l'élève à partir des synthèses des exercices précédents."""
    if not past_syntheses:
        return ""
    items = [s for s in past_syntheses if s.get("synthesis")]
    if not items:
        return ""
    formatted = "\n\n".join(
        f"- {s.get('title', 'Exercice')} :\n{s.get('synthesis', '')}" for s in items
    )
    return (
        "CARNET DE SUIVI DE L'ÉLÈVE (bilans des exercices précédents, du plus ancien au plus récent) :\n"
        f"{formatted}\n"
    )


@app.post("/classroom/exercise")
async def classroom_exercise(request: ClassroomExerciseRequest):
    """Le professeur propose 3 exercices d'écriture différents, parmi lesquels l'élève choisit."""
    logger.info(
        "[AI] route_started route=classroom_exercise model=%s previous_titles=%d",
        MISTRAL_BEST_MODEL,
        len(request.previous_exercise_titles),
    )
    avoid_repeats = ""
    if request.previous_exercise_titles:
        titles = "\n".join(f"- {t}" for t in request.previous_exercise_titles[-15:])
        avoid_repeats = (
            "L'élève a déjà fait les exercices suivants récemment, PROPOSE QUELQUE CHOSE DE DIFFÉRENT "
            f"(autre thème, autre type d'exercice, autre genre) :\n{titles}"
        )
    student_record = _build_student_record(request.past_syntheses)
    experience_context = _build_experience_context(request.experience_level)
    system_prompt = EXERCISE_GENERATION_PROMPT.format(
        teacher_name=TEACHER_NAME,
        teacher_personality=TEACHER_PROMPT,
        experience_context=experience_context,
        avoid_repeats=avoid_repeats,
        student_record=student_record,
    )
    try:
        raw = await generate_text_from_mistral(system_prompt, "Propose-moi 3 exercices au choix pour la classe.")
    except Exception as e:
        logger.exception("[AI] route_failed route=classroom_exercise model=%s", MISTRAL_BEST_MODEL)
        return {"error": str(e)}
    parts = re.split(r"=+\s*EXERCICE_SUIVANT\s*=+", raw)
    exercises = [p.strip() for p in parts if p.strip()]
    if not exercises:
        exercises = [raw.strip()]
    authors = []
    for ex in exercises:
        genre_match = re.search(r"\*\*Genre\s*:\*\*\s*(.+)", ex, re.IGNORECASE)
        author_id = get_author_for_genre(genre_match.group(1) if genre_match else None)
        authors.append({"id": author_id, "name": AUTHORS[author_id]["name"]})
    logger.info(
        "[AI] route_completed route=classroom_exercise model=%s exercises=%d response_chars=%d",
        MISTRAL_BEST_MODEL,
        len(exercises),
        len(raw),
    )
    return {"exercises": exercises, "authors": authors}


@app.post("/classroom/lesson")
async def classroom_lesson(request: ClassroomLessonRequest):
    """Le professeur donne un petit cours (bases + exemples) pour préparer l'élève à un exercice précis."""
    system_prompt = LESSON_PROMPT.format(
        teacher_name=TEACHER_NAME,
        teacher_personality=TEACHER_PROMPT,
        experience_context=_build_experience_context(request.experience_level),
        exercise_prompt=request.exercise_prompt,
    )
    return StreamingResponse(
        stream_from_mistral_small(system_prompt, "Donne-moi le cours pour me préparer à cet exercice."),
        media_type="text/plain"
    )


@app.post("/classroom/author-judgment")
async def classroom_author_judgment(request: ClassroomAuthorJudgmentRequest):
    """Un grand auteur (réel, passé ou présent) juge le travail de l'élève avec sa vraie personnalité."""
    author = AUTHORS.get(request.author_id)
    if not author:
        async def error_stream():
            yield "Auteur inconnu."
        return StreamingResponse(error_stream(), media_type="text/plain")

    peer_context = ""
    if request.peer_comments:
        peer_context += f"AVIS DES CAMARADES DE CLASSE :\n{request.peer_comments}\n\n"
    if request.teacher_critique:
        peer_context += f"AVIS DU PROFESSEUR {TEACHER_NAME} :\n{request.teacher_critique}\n\n"

    system_prompt = AUTHOR_JUDGMENT_PROMPT.format(
        author_name=author["name"],
        author_era=author["era"],
        teacher_name=TEACHER_NAME,
        exercise_prompt=request.exercise_prompt,
        experience_context=_build_experience_context(request.experience_level),
        author_personality=author["personality"],
        peer_context=peer_context,
    )
    # L'auteur est le rôle le plus exigeant en fidélité stylistique : on préfère Opus, puis GPT, puis Mistral.
    stream_fn, model, provider = _pick_provider([
        ("ANTHROPIC_API_KEY", stream_critique_from_anthropic, ANTHROPIC_BEST_MODEL),
        ("OPENAI_API_KEY", stream_critique_from_openai, OPENAI_BEST_MODEL),
        ("MISTRAL_API_KEY", stream_critique_from_mistral, MISTRAL_BEST_MODEL),
    ])
    return StreamingResponse(
        _trace_ai_response(stream_fn(request.text, system_prompt, model=model), provider, model),
        media_type="text/plain"
    )


@app.post("/classroom/peer-review")
async def classroom_peer_review(request: ClassroomPeerReviewRequest):
    """Un groupe d'élèves commente le travail rendu par l'utilisateur."""
    import random

    if request.student_ids:
        chosen = [sid for sid in request.student_ids if sid in STUDENTS]
    else:
        n = min(max(request.num_students, 1), len(STUDENT_IDS))
        chosen = random.sample(STUDENT_IDS, n)

    if not chosen:
        async def error_stream():
            yield "Aucun élève disponible pour commenter."
        return StreamingResponse(error_stream(), media_type="text/plain")

    students_list = "\n".join(f"- {STUDENTS[sid]['name']} ({STUDENTS[sid]['age']} ans)" for sid in chosen)
    students_personalities = "\n\n".join(
        f"{STUDENTS[sid]['name']} :\n{STUDENTS[sid]['personality']}" for sid in chosen
    )

    panel_system = STUDENTS_PANEL_PROMPT.format(
        nb_students=len(chosen),
        exercise_prompt=request.exercise_prompt,
        experience_context=_build_experience_context(request.experience_level),
        students_list=students_list,
        students_personalities=students_personalities,
    )

    chosen_names = [STUDENTS[sid]["name"] for sid in chosen]

    async def stream_with_header():
        yield f"*🧑‍🎓 {', '.join(chosen_names)} lisent votre texte…*\n\n---\n\n"
        # Le panel d'élèves n'a pas besoin du modèle le plus coûteux : dialogues courts et casual.
        async for chunk in _trace_ai_response(
            stream_critique_from_mistral(request.text, panel_system, model=MISTRAL_MEDIUM_MODEL),
            "mistral",
            MISTRAL_MEDIUM_MODEL,
        ):
            yield chunk

    return StreamingResponse(
        stream_with_header(),
        media_type="text/plain",
        headers={"X-Panel-Students": ",".join(chosen)}
    )


@app.post("/classroom/teacher-critique")
async def classroom_teacher_critique(request: ClassroomTeacherRequest):
    """Le professeur critique le travail de l'élève (bilan final ou aide à la demande)."""
    revision_note = TEACHER_REVISION_NOTE if request.text_changed_since_peer_review else ""
    if request.on_demand:
        peer_context = (
            f"Voici ce que les camarades ont déjà dit, si cela peut t'éclairer :\n{request.peer_comments}\n\n"
            if request.peer_comments else ""
        )
        mode_instructions = TEACHER_MODE_ON_DEMAND.format(peer_context=peer_context, revision_note=revision_note)
    else:
        mode_instructions = TEACHER_MODE_FINAL.format(
            peer_comments=request.peer_comments or "(l'élève n'a pas sollicité l'avis de ses camarades)",
            revision_note=revision_note,
        )

    student_record = _build_student_record(request.past_syntheses)
    system_prompt = TEACHER_CRITIQUE_PROMPT.format(
        teacher_name=TEACHER_NAME,
        teacher_personality=TEACHER_PROMPT,
        experience_context=_build_experience_context(request.experience_level),
        exercise_prompt=request.exercise_prompt,
        mode_instructions=mode_instructions,
        student_record=student_record,
    )

    stream_fn, model, provider = _pick_provider([
        ("ANTHROPIC_API_KEY", stream_critique_from_anthropic, ANTHROPIC_MEDIUM_MODEL),
        ("OPENAI_API_KEY", stream_critique_from_openai, OPENAI_MEDIUM_MODEL),
        ("MISTRAL_API_KEY", stream_critique_from_mistral, MISTRAL_MEDIUM_MODEL),
    ])
    return StreamingResponse(
        _trace_ai_response(stream_fn(request.text, system_prompt, model=model), provider, model),
        media_type="text/plain"
    )


@app.post("/classroom/synthesize")
async def classroom_synthesize(request: ClassroomSynthesisRequest):
    """Génère la synthèse +/- de fin d'exercice à conserver dans le carnet de suivi de l'élève."""
    system_prompt = SYNTHESIS_PROMPT.format(
        teacher_name=TEACHER_NAME,
        teacher_personality=TEACHER_PROMPT,
        experience_context=_build_experience_context(request.experience_level),
        exercise_prompt=request.exercise_prompt,
        critique=request.critique,
    )
    try:
        synthesis = await generate_text_from_mistral(system_prompt, "Rédige la synthèse pour le carnet de suivi.")
    except Exception as e:
        return {"error": str(e)}
    return {"synthesis": synthesis}


if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="0.0.0.0", port=8000)